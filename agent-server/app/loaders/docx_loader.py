"""
DOCX 加载器：按正文顺序遍历段落与表格（python-docx）。
关键修复：表格不再被拆成孤立单元格——2 列 / 偶数列按「标签-值」配对，
记录型表头（家庭成员等）按行输出。
"""
from __future__ import annotations

import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .document import Block, NormalizedDocument, render_for_llm
from .table_pairing import looks_record_header, pair_cells, strip_colon
from ..rules.patterns import match_section_title


def _cell_text(cell) -> str:
    parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    return " ".join(parts)


def _table_blocks(tbl_el, doc) -> list[Block]:
    table = Table(tbl_el, doc)
    blocks: list[Block] = []
    record_mode = False

    for row in table.rows:
        cells = [_cell_text(c) for c in row.cells]
        nonempty = [c for c in cells if c]
        if not nonempty:
            continue
        if not record_mode and looks_record_header(nonempty):
            record_mode = True
            continue
        if record_mode:
            blocks.append(Block(kind="record", text=" | ".join(nonempty)))
            continue
        pairs = pair_cells(cells)
        if pairs:
            for label, value in pairs:
                blocks.append(Block(kind="kv", label=label, value=value))
        else:
            blocks.extend(Block(kind="paragraph", text=c) for c in nonempty)

    # 嵌套表格（python-docx 单元格内的嵌套表，输出顺序靠后，不影响主流程）
    for nested in tbl_el.iter(qn("w:tbl")):
        if nested is tbl_el:
            continue
        blocks.extend(_table_blocks(nested, doc))
    return blocks


def _heading_level(paragraph: Paragraph) -> int:
    name = paragraph.style.name if paragraph.style else ""
    if name.startswith("Heading"):
        try:
            return int(name.split()[-1])
        except ValueError:
            return 1
    pPr = paragraph._p.pPr
    if pPr is not None and pPr.find(qn("w:outlineLvl")) is not None:
        return 1
    return 0


def load_docx(file_bytes: bytes, file_name: str) -> NormalizedDocument:
    import io

    doc = Document(io.BytesIO(file_bytes))
    blocks: list[Block] = []

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            if _heading_level(p) or match_section_title(text):
                blocks.append(Block(kind="heading", text=text))
            else:
                blocks.append(Block(kind="paragraph", text=text))
        elif child.tag == qn("w:tbl"):
            blocks.extend(_table_blocks(child, doc))

    cp = doc.core_properties
    metadata = {
        "author": (cp.author or "").strip(),
        "title": (cp.title or "").strip(),
        "comments": (cp.comments or "").strip(),
    }
    return NormalizedDocument(
        file_name=file_name,
        file_type="docx",
        blocks=blocks,
        text=render_for_llm(blocks),
        metadata=metadata,
    )
